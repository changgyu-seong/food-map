import random

from flask import Flask, render_template, request
from openpyxl import load_workbook


app = Flask(__name__)
TARGET_SHEET_NAME = "서현"
HEADER_KEYWORDS = ("식당", "가게", "음식점", "상호", "이름", "매장", "업체")


def extract_restaurants(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header_row = rows[0]
    restaurant_col_index = None

    for index, value in enumerate(header_row):
        if isinstance(value, str) and any(keyword in value for keyword in HEADER_KEYWORDS):
            restaurant_col_index = index
            break

    data_rows = rows[1:] if restaurant_col_index is not None else rows
    restaurants = []

    for row in data_rows:
        if restaurant_col_index is not None:
            value = row[restaurant_col_index] if restaurant_col_index < len(row) else None
        else:
            value = next((cell for cell in row if isinstance(cell, str) and cell.strip()), None)

        if isinstance(value, str):
            name = value.strip()
            if name and name not in HEADER_KEYWORDS:
                restaurants.append(name)

    return restaurants


def choose_random_restaurant(file_storage):
    if file_storage is None or not file_storage.filename:
        raise FileNotFoundError("업로드할 .xlsx 파일을 선택해 주세요.")

    if not file_storage.filename.lower().endswith(".xlsx"):
        raise ValueError(".xlsx 파일만 업로드할 수 있습니다.")

    workbook = load_workbook(file_storage.stream, data_only=True)

    if TARGET_SHEET_NAME not in workbook.sheetnames:
        raise KeyError(f"'{TARGET_SHEET_NAME}' 시트를 찾을 수 없습니다.")

    restaurants = extract_restaurants(workbook[TARGET_SHEET_NAME])
    if not restaurants:
        raise ValueError("선택할 수 있는 식당 데이터가 없습니다.")

    return file_storage.filename, restaurants, random.choice(restaurants)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/select", methods=["GET", "POST"])
def select_restaurant():
    if request.method == "GET":
        return render_template(
            "select.html",
            restaurant=None,
            restaurants=[],
            workbook_name=None,
            sheet_name=TARGET_SHEET_NAME,
            error=None,
        )

    try:
        workbook_name, restaurants, restaurant = choose_random_restaurant(
            request.files.get("file")
        )
        return render_template(
            "select.html",
            restaurant=restaurant,
            restaurants=restaurants,
            workbook_name=workbook_name,
            sheet_name=TARGET_SHEET_NAME,
            error=None,
        )
    except (FileNotFoundError, KeyError, ValueError) as error:
        return render_template(
            "select.html",
            restaurant=None,
            restaurants=[],
            workbook_name=None,
            sheet_name=TARGET_SHEET_NAME,
            error=str(error),
        ), 400


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
